"""Gera fichas_maquinas.json, chinaglass.xlsx e o catálogo em Markdown a partir de dados/."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
from fichas import COLUNAS, FICHAS, PECAS

OUT = sys.argv[1]
D = os.path.join(OUT, "dados")
load = lambda n: json.load(open(os.path.join(D, n), encoding="utf-8"))
produtos, categorias, site = load("produtos.json"), load("categorias.json"), load("site.json")
por_id = {p["id"]: p for p in produtos}

pasta = lambda p: os.path.dirname(p["imagens_galeria"][0]) if p["imagens_galeria"] else ""
base = lambda p: {
    "id": p["id"], "codigo": p["sku"], "nome": p["nome"],
    "categorias": ", ".join(p["categorias"]) or ", ".join(p["trilha"][1:]),
    "url": p["url"], "video": " ".join(p["videos_youtube"]),
    "fotos": len(p["imagens_galeria"]), "imagens_descricao": len(p["imagens_descricao"]),
    "pasta_imagens": pasta(p),
}

maquinas = []
for pid, f in FICHAS.items():
    p = por_id[pid]
    row = base(p) | {k: f.get(k, "") for k, _ in COLUNAS}
    row["spec_em_imagem"] = "Sim" if f.get("spec_em_imagem") else ""
    maquinas.append(row)
maquinas.sort(key=lambda r: (r["familia"], r["nome"]))
pecas = [base(p) | {"detalhes": PECAS.get(p["id"], p["descricao"])} for p in produtos if p["id"] not in FICHAS]
assert len(maquinas) + len(pecas) == len(produtos)

json.dump({"maquinas": maquinas, "pecas": pecas}, open(os.path.join(D, "fichas_maquinas.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)

# ---------- Excel ----------
wb = Workbook()
HEAD = PatternFill("solid", fgColor="1F2937")


def sheet(ws, header, rows, widths=None):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink, c.style = c.value.split(" ")[0], "Hyperlink"
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


ws = wb.active
ws.title = "Máquinas"
cols = [("nome", "Máquina", 45), ("codigo", "Código", 14), ("categorias", "Categorias no site", 22)] + \
       [(k, t, 30 if k in ("comando", "marcas") else 50 if k == "destaques" else 18) for k, t in COLUNAS] + \
       [("spec_em_imagem", "Ficha técnica só em imagem", 12), ("fotos", "Nº fotos", 8), ("video", "Vídeo", 30),
        ("url", "Página", 40), ("pasta_imagens", "Pasta das imagens", 40)]
sheet(ws, [c[1] for c in cols], [[m[c[0]] for c in cols] for m in maquinas], [c[2] for c in cols])

ws = wb.create_sheet("Peças")
sheet(ws, ["Peça", "Código", "Detalhes", "Nº fotos", "Página", "Pasta das imagens"],
      [[p["nome"], p["codigo"], p["detalhes"], p["fotos"], p["url"], p["pasta_imagens"]] for p in pecas],
      [40, 22, 50, 8, 45, 45])

ws = wb.create_sheet("Categorias")
sheet(ws, ["Categoria", "Nº produtos", "Produtos", "Página", "Título da página", "Meta description"],
      [[c["nome"], len(c["produtos"]), "\n".join(c["produtos"]), c["url"], c["titulo_pagina"], c["meta_description"]] for c in categorias],
      [28, 10, 70, 40, 40, 60])

ws = wb.create_sheet("Descrições completas")
sheet(ws, ["Produto", "Código", "Título SEO", "Meta description", "Descrição (texto do site)"],
      [[p["nome"], p["sku"], p["titulo_pagina"], p["meta_description"], p["descricao"][:32000]] for p in produtos],
      [40, 14, 40, 50, 120])

ws = wb.create_sheet("Imagens")
rows = []
for p in produtos:
    for tipo, arqs, urls in (("Galeria", p["imagens_galeria"], p["imagens_galeria_url"]),
                             ("Descrição", p["imagens_descricao"], p["imagens_descricao_url"])):
        rows += [[p["nome"], tipo, a, u] for a, u in zip(arqs, urls)]
rows += [["(Home) " + (b["alt"] or b["tipo"]), b["tipo"].capitalize(), b["arquivo"], b["url"]] for b in site["banners_e_logos"]]
sheet(ws, ["Produto / origem", "Tipo", "Arquivo local", "URL original"], rows, [45, 12, 70, 80])

ws = wb.create_sheet("Empresa")
info = [("Nome", site["nome"]), ("Site", site["url"]), ("Plataforma", site["plataforma"]), ("CNPJ", site["cnpj"]),
        ("Telefone", ", ".join(site["telefones"])), ("E-mail", ", ".join(site["emails"])),
        ("Descrição", site["meta_description"]), ("Sobre", site["sobre"]),
        ("Redes / links externos", "\n".join(site["redes_sociais"])),
        ("Totais", f"{len(maquinas)} máquinas, {len(pecas)} peças, {len(categorias)} categorias, {len(rows)} imagens")]
sheet(ws, ["Campo", "Valor"], info, [22, 100])
wb.save(os.path.join(OUT, "chinaglass.xlsx"))

# ---------- Catálogo Markdown ----------
L = [f"# Catálogo China Glass — máquinas\n", f"Fonte: {site['url']} · {len(maquinas)} máquinas · {len(pecas)} peças de reposição\n"]
fam = None
for m in maquinas:
    L.append(f"\n## {m['nome']}\n")
    L.append(f"- **Código:** {m['codigo']} · **Categoria:** {m['categorias']}")
    for k, t in COLUNAS:
        if m[k]:
            L.append(f"- **{t}:** {m[k]}")
    if m["spec_em_imagem"]:
        L.append("- **Obs.:** a ficha técnica completa está em imagem (pasta `descricao/`).")
    L.append(f"- **Fotos:** {m['fotos']} em `{m['pasta_imagens']}`")
    if m["video"]:
        L.append(f"- **Vídeo:** {m['video']}")
    L.append(f"- **Página:** {m['url']}")
L.append("\n## Peças de reposição\n")
L += [f"- {p['nome']} ({p['codigo']}) — {p['detalhes']}" for p in pecas]
open(os.path.join(OUT, "catalogo_maquinas.md"), "w", encoding="utf-8").write("\n".join(L) + "\n")
print("maquinas", len(maquinas), "pecas", len(pecas), "imagens", len(rows))
